import os
import uuid
import logging
from fastapi import UploadFile
from typing import Dict, Any
import PyPDF2
import docx
import csv
import json
from app.core.config import settings
from app.models.document import Document

logger = logging.getLogger(__name__)

# Asegurar que el directorio de uploads existe
os.makedirs(settings.UPLOAD_DIR, exist_ok=True)

async def process_document(file: UploadFile, document_type: str) -> Document:
    """
    Procesa un documento subido, extrae su contenido y lo guarda
    """
    # Generar un ID único para el documento
    document_id = str(uuid.uuid4())
    
    # Crear directorio para este documento
    document_dir = os.path.join(settings.UPLOAD_DIR, document_id)
    os.makedirs(document_dir, exist_ok=True)
    
    # Guardar el archivo original
    file_path = os.path.join(document_dir, file.filename)
    
    try:
        # Guardar el archivo
        with open(file_path, "wb") as f:
            content = await file.read()
            f.write(content)
        
        # Extraer texto según el tipo de documento
        if document_type == "pdf":
            text_content = extract_text_from_pdf(file_path)
        elif document_type == "docx":
            text_content = extract_text_from_docx(file_path)
        elif document_type == "txt":
            text_content = extract_text_from_txt(file_path)
        elif document_type == "csv":
            text_content = extract_text_from_csv(file_path)
        elif document_type == "json":
            text_content = extract_text_from_json(file_path)
        else:
            raise ValueError(f"Tipo de documento no soportado: {document_type}")
        
        # Guardar el texto extraído
        text_path = os.path.join(document_dir, "extracted_text.txt")
        with open(text_path, "w", encoding="utf-8") as f:
            f.write(text_content)
        
        # Crear y devolver el objeto documento
        document = Document(
            id=document_id,
            filename=file.filename,
            document_type=document_type,
            file_path=file_path,
            text_path=text_path,
            text_content=text_content[:1000] + "..." if len(text_content) > 1000 else text_content,  # Versión truncada para la respuesta
            size_bytes=os.path.getsize(file_path),
            metadata={
                "upload_timestamp": str(os.path.getctime(file_path)),
                "content_length": len(text_content)
            }
        )
        
        return document
        
    except Exception as e:
        # Limpiar en caso de error
        if os.path.exists(document_dir):
            import shutil
            shutil.rmtree(document_dir)
        
        logger.error(f"Error procesando documento: {str(e)}")
        raise ValueError(f"Error procesando documento: {str(e)}")

def extract_text_from_pdf(file_path: str) -> str:
    """Extrae texto de un archivo PDF"""
    text = ""
    try:
        with open(file_path, "rb") as f:
            pdf_reader = PyPDF2.PdfReader(f)
            for page_num in range(len(pdf_reader.pages)):
                text += pdf_reader.pages[page_num].extract_text() + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extrayendo texto de PDF: {str(e)}")
        raise ValueError(f"Error extrayendo texto de PDF: {str(e)}")

def extract_text_from_docx(file_path: str) -> str:
    """Extrae texto de un archivo DOCX"""
    try:
        doc = docx.Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except Exception as e:
        logger.error(f"Error extrayendo texto de DOCX: {str(e)}")
        raise ValueError(f"Error extrayendo texto de DOCX: {str(e)}")

def extract_text_from_txt(file_path: str) -> str:
    """Extrae texto de un archivo TXT"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        # Intentar con otra codificación si utf-8 falla
        with open(file_path, "r", encoding="latin-1") as f:
            return f.read()
    except Exception as e:
        logger.error(f"Error extrayendo texto de TXT: {str(e)}")
        raise ValueError(f"Error extrayendo texto de TXT: {str(e)}")

def extract_text_from_csv(file_path: str) -> str:
    """Convierte un CSV a texto estructurado"""
    try:
        text = ""
        with open(file_path, "r", encoding="utf-8") as f:
            csv_reader = csv.reader(f)
            for row in csv_reader:
                text += " | ".join(row) + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extrayendo texto de CSV: {str(e)}")
        raise ValueError(f"Error extrayendo texto de CSV: {str(e)}")

def extract_text_from_json(file_path: str) -> str:
    """Extrae texto de un archivo JSON"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        # Convertir el JSON a texto formateado
        return json.dumps(data, indent=2, ensure_ascii=False)
    except Exception as e:
        logger.error(f"Error extrayendo texto de JSON: {str(e)}")
        raise ValueError(f"Error extrayendo texto de JSON: {str(e)}")